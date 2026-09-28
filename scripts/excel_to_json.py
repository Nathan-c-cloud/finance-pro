#!/usr/bin/env python3
"""
excel_to_json.py — Convertit le fichier Excel "Suivi_Budget" (feuilles
Référentiel_Fixes, Soldes, Transactions) en un fichier JSON prêt à être
importé dans l'application Angular/Supabase via DataService.importFromJson.

Usage :
    python3 excel_to_json.py chemin/vers/Suivi_Budget.xlsm [sortie.json]

Le script ne modifie jamais le fichier Excel d'origine : il ne fait que le
lire. Le JSON produit peut ensuite être importé depuis la page "Réglages"
de l'application (bouton "Importer depuis Excel").

Hypothèses sur la structure du fichier (celle du fichier livré par Claude,
avec les tables Excel TableFixes / TableSoldes / Tableau1) :
  - Référentiel_Fixes : table TableFixes (Nom, Montant (€), Catégorie,
    Fréquence, Jour de prélèvement, Actif) + liste des catégories en
    colonne H (une catégorie par ligne, sert de source à la liste
    déroulante).
  - Soldes : table TableSoldes (Mois, ..., Statut, Solde début (forcer,
    optionnel)) + cellule K1 = solde initial.
  - Transactions : table Tableau1 (Mois, Type, Nom, Montant (€),
    Catégorie, Date, Détail, Nécessaire ?, Reçu ?) + cellule L6 = solde
    réel constaté (rapprochement bancaire, rattaché au mois affiché le
    plus récent).

Si le fichier n'a pas ces tables nommées (ex. fichier jamais réparé),
le script retombe sur des plages fixes approximatives — vérifie alors le
JSON produit avant de l'importer.
"""

import json
import sys
from datetime import datetime, date
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Ce script a besoin d'openpyxl : pip install --break-system-packages openpyxl")
    sys.exit(1)


def cell_str(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def to_date_str(v):
    """Excel peut renvoyer un datetime, une date, ou parfois une string."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    return s or None


def to_month_date_str(v):
    """Normalise une date de mois en 'YYYY-MM-01' (le jour n'a pas de sens ici)."""
    d = to_date_str(v)
    if not d:
        return None
    return d[:8] + "01"


def to_bool(v):
    """'Oui' -> True, 'Non' -> False, vide/None -> None (pas de valeur renseignée)."""
    if v is None:
        return None
    s = str(v).strip().lower()
    if s in ("oui", "yes", "true", "1", "x"):
        return True
    if s in ("non", "no", "false", "0"):
        return False
    return None


def table_range(ws, table_name, fallback_ref):
    tables = dict(ws.tables.items()) if hasattr(ws, "tables") else {}
    ref = tables.get(table_name, fallback_ref)
    return openpyxl.utils.cell.range_boundaries(ref)  # (min_col, min_row, max_col, max_row)


def read_categories(ws_fixes):
    cats = []
    seen = set()
    for r in range(1, 200):
        v = cell_str(ws_fixes.cell(row=r, column=8).value)
        if v and v not in seen:
            cats.append(v)
            seen.add(v)
    return cats, seen


def read_fixed_expenses(ws_fixes, known_categories, seen_categories):
    min_col, min_row, max_col, max_row = table_range(ws_fixes, "TableFixes", "A4:F21")
    header_row = min_row
    headers = [cell_str(ws_fixes.cell(row=header_row, column=c).value) for c in range(min_col, max_col + 1)]
    col_index = {h: i for i, h in enumerate(headers) if h}

    def get(row_vals, name, default=None):
        i = col_index.get(name)
        return row_vals[i] if i is not None and i < len(row_vals) else default

    out = []
    for r in range(header_row + 1, max_row + 1):
        row_vals = [ws_fixes.cell(row=r, column=c).value for c in range(min_col, max_col + 1)]
        name = cell_str(get(row_vals, "Nom"))
        amount = get(row_vals, "Montant (€)")
        if not name or amount is None:
            continue
        category = cell_str(get(row_vals, "Catégorie"))
        if category and category not in seen_categories:
            known_categories.append(category)
            seen_categories.add(category)
        frequency_raw = (cell_str(get(row_vals, "Fréquence")) or "Mensuel").lower()
        frequency = "annual" if frequency_raw.startswith("annu") else "monthly"
        payment_day = get(row_vals, "Jour de prélèvement")
        try:
            payment_day = int(payment_day) if payment_day not in (None, "") else 1
        except (ValueError, TypeError):
            payment_day = 1
        active = to_bool(get(row_vals, "Actif"))
        out.append(
            {
                "name": name,
                "amount": round(float(amount), 2),
                "category": category,
                "frequency": frequency,
                "paymentDay": max(1, min(28, payment_day)),
                "active": active if active is not None else True,
            }
        )
    return out


def read_months(ws_soldes):
    min_col, min_row, max_col, max_row = table_range(ws_soldes, "TableSoldes", "A4:I6")
    header_row = min_row
    headers = [cell_str(ws_soldes.cell(row=header_row, column=c).value) for c in range(min_col, max_col + 1)]
    col_index = {h: i for i, h in enumerate(headers) if h}

    def get(row_vals, name, default=None):
        i = col_index.get(name)
        return row_vals[i] if i is not None and i < len(row_vals) else default

    out = []
    for r in range(header_row + 1, max_row + 1):
        row_vals = [ws_soldes.cell(row=r, column=c).value for c in range(min_col, max_col + 1)]
        month_raw = get(row_vals, "Mois")
        month_date = to_month_date_str(month_raw)
        if not month_date:
            continue
        statut = (cell_str(get(row_vals, "Statut")) or "").lower()
        status = "closed" if "clôtur" in statut or "clotur" in statut else "current"
        override = get(row_vals, "Solde début\n(forcer, optionnel)")
        if override is None:
            # openpyxl peut aussi renvoyer l'en-tête sans le saut de ligne selon la version du fichier
            override = get(row_vals, "Solde début (forcer, optionnel)")
        out.append(
            {
                "monthDate": month_date,
                "status": status,
                "startingBalanceOverride": round(float(override), 2) if override not in (None, "") else None,
                "realBalanceCheck": None,  # complété après coup pour le mois le plus récent
            }
        )
    return out


def read_transactions(ws_tx, known_categories, seen_categories):
    min_col, min_row, max_col, max_row = table_range(ws_tx, "Tableau1", "A11:I99")
    header_row = min_row
    headers = [cell_str(ws_tx.cell(row=header_row, column=c).value) for c in range(min_col, max_col + 1)]
    col_index = {h: i for i, h in enumerate(headers) if h}

    def get(row_vals, name, default=None):
        i = col_index.get(name)
        return row_vals[i] if i is not None and i < len(row_vals) else default

    type_map = {"fixe": "fixed", "variable": "variable", "revenu": "income"}

    out = []
    for r in range(header_row + 1, max_row + 1):
        row_vals = [ws_tx.cell(row=r, column=c).value for c in range(min_col, max_col + 1)]
        month_raw = get(row_vals, "Mois")
        month_date = to_month_date_str(month_raw)
        name = cell_str(get(row_vals, "Nom"))
        if not month_date or not name:
            continue
        amount = get(row_vals, "Montant (€)")
        if amount is None:
            continue
        type_raw = (cell_str(get(row_vals, "Type")) or "variable").strip().lower()
        tx_type = type_map.get(type_raw, "variable")
        category = cell_str(get(row_vals, "Catégorie"))
        if category and category not in seen_categories:
            known_categories.append(category)
            seen_categories.add(category)
        out.append(
            {
                "monthDate": month_date,
                "type": tx_type,
                "name": name,
                "amount": round(abs(float(amount)), 2),
                "category": category,
                "txDate": to_date_str(get(row_vals, "Date")),
                "detail": cell_str(get(row_vals, "Détail")),
                "necessary": to_bool(get(row_vals, "Nécessaire ?")),
                "received": to_bool(get(row_vals, "Reçu ?")),
            }
        )
    return out


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    src = Path(sys.argv[1])
    if not src.exists():
        print(f"Fichier introuvable : {src}")
        sys.exit(1)

    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".import.json")

    wb = openpyxl.load_workbook(src, data_only=True, keep_vba=False)

    ws_fixes = wb["Référentiel_Fixes"]
    ws_soldes = wb["Soldes"]
    ws_tx = wb["Transactions"]

    categories, seen_categories = read_categories(ws_fixes)
    fixed_expenses = read_fixed_expenses(ws_fixes, categories, seen_categories)
    months = read_months(ws_soldes)
    transactions = read_transactions(ws_tx, categories, seen_categories)

    initial_balance = ws_soldes["K1"].value
    try:
        initial_balance = round(float(initial_balance), 2) if initial_balance is not None else 0
    except (ValueError, TypeError):
        initial_balance = 0

    # Rapprochement bancaire (Transactions!L6, à droite du libellé en K6) :
    # rattaché au mois le plus récent, celui affiché dans l'app au moment de l'export.
    real_balance = ws_tx["L6"].value
    if months and real_balance is not None:
        try:
            real_balance = round(float(real_balance), 2)
            latest = max(months, key=lambda m: m["monthDate"])
            latest["realBalanceCheck"] = real_balance
        except (ValueError, TypeError):
            pass

    payload = {
        "categories": categories,
        "fixedExpenses": fixed_expenses,
        "months": months,
        "transactions": transactions,
        "initialBalance": initial_balance,
    }

    dst.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"OK — {len(categories)} catégories, {len(fixed_expenses)} dépenses fixes, "
          f"{len(months)} mois, {len(transactions)} transactions.")
    print(f"Fichier écrit : {dst}")
    print("Importe ce fichier depuis la page Réglages de l'application (\"Importer depuis Excel\").")


if __name__ == "__main__":
    main()
